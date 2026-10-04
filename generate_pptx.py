import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

THEMES = {
    "Unit_1": {
        "primary": RGBColor(13, 71, 161),       # Deep Blue
        "accent": RGBColor(21, 101, 192),       # Medium Blue
        "light_bg": RGBColor(240, 244, 248),    # Soft blue-gray
        "title_bg": RGBColor(10, 37, 64),       # Navy Dark
        "badge_bg": RGBColor(30, 136, 229),     # Bright Blue
    },
    "Unit_2": {
        "primary": RGBColor(230, 81, 0),        # Deep Orange
        "accent": RGBColor(245, 124, 0),        # Orange Accent
        "light_bg": RGBColor(255, 243, 224),    # Soft cream orange
        "title_bg": RGBColor(62, 39, 35),       # Dark Umber
        "badge_bg": RGBColor(251, 140, 0),      # Bright Orange
    },
    "Unit_3": {
        "primary": RGBColor(49, 27, 146),       # Deep Purple/Indigo
        "accent": RGBColor(69, 39, 160),        # Violet
        "light_bg": RGBColor(237, 231, 246),    # Soft lilac
        "title_bg": RGBColor(26, 15, 60),       # Dark Purple
        "badge_bg": RGBColor(103, 58, 183),     # Bright Purple
    },
    "Unit_4": {
        "primary": RGBColor(0, 77, 64),         # Deep Teal
        "accent": RGBColor(0, 105, 92),         # Medium Teal
        "light_bg": RGBColor(224, 242, 241),    # Soft mint
        "title_bg": RGBColor(13, 43, 39),       # Dark Teal
        "badge_bg": RGBColor(0, 150, 136),      # Bright Teal
    }
}

def clean_text(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    text = re.sub(r'\$(.*?)\$', r'\1', text)
    text = text.replace('&le;', '<=').replace('&ge;', '>=').replace('&ne;', '!=').replace('&rarr;', '->')
    return text.strip()

def build_presentation(md_path, pptx_path, unit_key, division_name):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]
    
    theme = THEMES.get(unit_key, THEMES["Unit_1"])
    
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    # Split slides by standard markdown horizontal rule
    raw_slides = re.split(r'\n---\n|\n---[ \t]*\r?\n', content)
    
    slide_index = 0
    for raw in raw_slides:
        raw = raw.strip()
        if not raw or raw.startswith("marp: true") or "theme: gaia" in raw:
            continue
            
        slide = prs.slides.add_slide(blank_slide_layout)
        slide_index += 1
        
        # Check if Lead Slide (Class title banner or Master deck title)
        is_lead = "<!-- _class: lead -->" in raw or raw.startswith("# ")
        clean_raw = raw.replace("<!-- _class: lead -->", "").strip()
        
        lines = [l.rstrip() for l in clean_raw.splitlines() if l.strip()]
        if not lines:
            continue

        if is_lead and (lines[0].startswith("# ") or lines[0].startswith("## ")):
            # Render Lead Banner Slide
            # Background shape
            bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
            bg_shape.fill.solid()
            bg_shape.fill.fore_color.rgb = theme["title_bg"]
            bg_shape.line.fill.background()

            # Accent color bar on left
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.4), Inches(7.5))
            bar.fill.solid()
            bar.fill.fore_color.rgb = theme["badge_bg"]
            bar.line.fill.background()

            # Badge pill
            badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.2), Inches(3.8), Inches(0.6))
            badge.fill.solid()
            badge.fill.fore_color.rgb = theme["badge_bg"]
            badge.line.fill.background()
            badge_tf = badge.text_frame
            badge_tf.text = f"ACSL • {unit_key.replace('_', ' ')} • {division_name.upper()}"
            badge_tf.paragraphs[0].font.size = Pt(13)
            badge_tf.paragraphs[0].font.bold = True
            badge_tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
            badge_tf.paragraphs[0].alignment = PP_ALIGN.CENTER

            # Main Title & Subtitle box
            tb = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11.0), Inches(4.2))
            tf = tb.text_frame
            tf.word_wrap = True

            first_p = True
            for l in lines:
                l_strip = l.strip()
                p = tf.paragraphs[0] if first_p else tf.add_paragraph()
                first_p = False

                if l_strip.startswith("# "):
                    p.text = clean_text(l_strip[2:])
                    p.font.size = Pt(36)
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(255, 255, 255)
                    p.space_after = Pt(14)
                elif l_strip.startswith("## "):
                    p.text = clean_text(l_strip[3:])
                    p.font.size = Pt(24)
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(200, 225, 255)
                    p.space_after = Pt(10)
                elif l_strip.startswith("### "):
                    p.text = clean_text(l_strip[4:])
                    p.font.size = Pt(18)
                    p.font.color.rgb = RGBColor(176, 190, 205)
                    p.space_after = Pt(8)
                else:
                    p.text = clean_text(l_strip)
                    p.font.size = Pt(16)
                    p.font.color.rgb = RGBColor(230, 235, 245)
                    p.space_after = Pt(6)

        else:
            # Standard Content Slide
            # Top header bar
            header_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.2))
            header_bar.fill.solid()
            header_bar.fill.fore_color.rgb = theme["primary"]
            header_bar.line.fill.background()

            # Unit tag in header
            tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.3))
            tag_tf = tag_box.text_frame
            tag_p = tag_tf.paragraphs[0]
            tag_p.text = f"{unit_key.replace('_', ' ')} | {division_name.capitalize()} Division"
            tag_p.font.size = Pt(10)
            tag_p.font.color.rgb = RGBColor(187, 222, 251)
            tag_p.font.bold = True

            # Extract slide title
            title_text = "Key Concept"
            body_lines = []
            for l in lines:
                l_strip = l.strip()
                if l_strip.startswith("## "):
                    title_text = clean_text(l_strip[3:])
                elif l_strip.startswith("# "):
                    title_text = clean_text(l_strip[2:])
                else:
                    body_lines.append(l)

            # Slide Title Text
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.5), Inches(0.7))
            title_tf = title_box.text_frame
            title_p = title_tf.paragraphs[0]
            title_p.text = title_text
            title_p.font.size = Pt(22)
            title_p.font.bold = True
            title_p.font.color.rgb = RGBColor(255, 255, 255)

            # Check if there is a code block or table
            raw_body = "\n".join(body_lines)
            has_code = "```" in raw_body
            has_table = "|" in raw_body and "-|-" in raw_body or "| :-" in raw_body

            # Content container
            content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3))
            tf = content_box.text_frame
            tf.word_wrap = True

            in_code = False
            code_buffer = []
            first_p = True

            for bl in body_lines:
                bl_strip = bl.strip()
                if bl_strip.startswith("```"):
                    if in_code:
                        # Flush code card
                        in_code = False
                        p = tf.paragraphs[0] if first_p else tf.add_paragraph()
                        first_p = False
                        p.text = "\n".join(code_buffer)
                        p.font.name = "Consolas"
                        p.font.size = Pt(13)
                        p.font.color.rgb = RGBColor(33, 33, 33)
                        p.space_after = Pt(10)
                        code_buffer = []
                    else:
                        in_code = True
                        code_buffer = []
                    continue

                if in_code:
                    code_buffer.append(bl)
                    continue

                if bl_strip.startswith("|") and ("---" in bl_strip or ":-" in bl_strip):
                    continue

                if not bl_strip:
                    continue

                p = tf.paragraphs[0] if first_p else tf.add_paragraph()
                first_p = False

                if bl_strip.startswith("### "):
                    p.text = clean_text(bl_strip[4:])
                    p.font.size = Pt(17)
                    p.font.bold = True
                    p.font.color.rgb = theme["accent"]
                    p.space_before = Pt(8)
                    p.space_after = Pt(4)
                elif bl_strip.startswith("- ") or bl_strip.startswith("* "):
                    p.text = "•  " + clean_text(bl_strip[2:])
                    p.font.size = Pt(15)
                    p.font.color.rgb = RGBColor(38, 50, 56)
                    p.space_after = Pt(4)
                    p.level = 0
                elif re.match(r'^\d+\.\s', bl_strip):
                    p.text = clean_text(bl_strip)
                    p.font.size = Pt(15)
                    p.font.color.rgb = RGBColor(38, 50, 56)
                    p.space_after = Pt(4)
                elif bl_strip.startswith("|"):
                    # Table row representation
                    cells = [clean_text(c) for c in bl_strip.split("|")[1:-1]]
                    p.text = "    ".join(cells)
                    p.font.name = "Consolas"
                    p.font.size = Pt(13)
                    p.font.color.rgb = RGBColor(55, 71, 79)
                    p.space_after = Pt(2)
                else:
                    p.text = clean_text(bl_strip)
                    p.font.size = Pt(15)
                    p.font.color.rgb = RGBColor(38, 50, 56)
                    p.space_after = Pt(6)

            # Footer
            footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.4))
            footer_tf = footer_box.text_frame
            footer_p = footer_tf.paragraphs[0]
            footer_p.text = f"ACSL Tutoring Program • {division_name.capitalize()} Division • Page {slide_index}"
            footer_p.font.size = Pt(10)
            footer_p.font.color.rgb = RGBColor(144, 164, 174)

    prs.save(pptx_path)
    print(f"Generated: {pptx_path} ({slide_index} slides)")

if __name__ == "__main__":
    base_dir = r"c:\Users\javaone\Documents\ACLC"
    units = [
        "Unit_1_Number_Systems",
        "Unit_2_Bit_String_and_Notation",
        "Unit_3_Boolean_Algebra_and_Data_Structures",
        "Unit_4_Graph_Theory_and_Digital_Electronics"
    ]
    divisions = ["elementary", "junior", "intermediate"]

    for unit in units:
        unit_key = unit.split("_")[0] + "_" + unit.split("_")[1] # e.g. Unit_1
        for div in divisions:
            md_file = os.path.join(base_dir, unit, div, "slides.md")
            pptx_file = os.path.join(base_dir, unit, div, "slides.pptx")
            if os.path.exists(md_file):
                build_presentation(md_file, pptx_file, unit_key, div)
